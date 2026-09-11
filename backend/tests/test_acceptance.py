import io
import json
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from main import app
from database import Base, engine, SessionLocal
from models import Job, Candidate, Match, Analysis

client = TestClient(app)


def setup_function():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def make_docx():
    b = io.BytesIO(); doc = Document(); doc.add_paragraph("Alice Example"); doc.add_paragraph("alice@example.com"); doc.add_paragraph("Python FastAPI SQL"); doc.add_paragraph("3 years experience"); doc.add_paragraph("Bachelor's degree"); doc.save(b); b.seek(0); return b


def make_pdf():
    b = io.BytesIO(); c = canvas.Canvas(b); c.drawString(72, 750, "Bob Example"); c.drawString(72, 730, "bob@example.com"); c.drawString(72, 710, "Python FastAPI SQL"); c.drawString(72, 690, "2 years experience"); c.drawString(72, 670, "Bachelor's degree"); c.save(); b.seek(0); return b


def test_document_parsing_and_upload():
    r = client.post("/api/resumes/upload", files=[("files", ("alice.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")), ("files", ("bob.pdf", make_pdf(), "application/pdf"))])
    assert r.status_code == 200
    data = r.json()["data"]
    assert len(data) == 2
    assert {x["resume_filename"] for x in data} == {"alice.docx", "bob.pdf"}


def test_invalid_upload_returns_consistent_error():
    r = client.post("/api/resumes/upload", files=[("files", ("bad.exe", b"not a document", "application/octet-stream"))])
    assert r.status_code == 400
    assert r.json()["success"] is False
    assert "error" in r.json()


def test_three_jobs_ten_candidates_full_pipeline():
    # Seed through the public data layer to emulate the required 3-job/10-candidate demo workflow.
    db = SessionLocal()
    for j in range(3):
        db.add(Job(title=f"Synthetic Job {j+1}", description="Python FastAPI SQL, 2 years experience, bachelor's degree.", required_skills=json.dumps(["python", "fastapi", "sql"]), preferred_skills=json.dumps(["docker"]), minimum_experience=2, education_requirements="bachelor", other_requirements="[]", extracted_requirements="{}"))
    for i in range(10):
        db.add(Candidate(name=f"Synthetic Candidate {i+1}", email=f"c{i+1}@example.com", resume_filename=f"c{i+1}.docx", resume_path=f"uploads/resumes/c{i+1}.docx", raw_resume_text="Python FastAPI SQL", skills=json.dumps(["python", "fastapi", "sql"]), education="Bachelor's degree", experience=2 + (i % 3)))
    db.commit(); db.close()
    for job_id in [1, 2, 3]:
        r = client.post(f"/api/matching/run/{job_id}")
        assert r.status_code == 200
        assert len(r.json()["data"]["results"]) == 10
    db = SessionLocal()
    assert db.query(Job).count() == 3
    assert db.query(Candidate).count() == 10
    assert db.query(Match).count() == 30
    assert db.query(Analysis).count() == 30
    db.close()
