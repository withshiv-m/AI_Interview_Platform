from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


class DocumentParseError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(file_path):
    try:
        reader = PdfReader(str(file_path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentParseError(f"Could not parse PDF: {exc}") from exc
    text = _clean(text)
    if not text:
        raise DocumentParseError("The PDF contains no extractable text.")
    return text


def extract_text_from_docx(file_path):
    try:
        doc = Document(str(file_path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    except Exception as exc:
        raise DocumentParseError(f"Could not parse DOCX: {exc}") from exc
    text = _clean(text)
    if not text:
        raise DocumentParseError("The DOCX contains no extractable text.")
    return text


def extract_text(file_path):
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    raise DocumentParseError("Unsupported document type. Only PDF and DOCX are allowed.")
